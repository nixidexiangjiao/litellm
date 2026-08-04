"""The model price list.

Prices live in DuckDB (``model_prices``) so that reporting is plain SQL and a
corrected price re-values past runs. This module only owns the shape of a price
row and the optional Excel importer used to seed the table; the arithmetic is
in ``storage.SCHEMA``.

``openpyxl`` is imported inside the importer so the row shape stays usable
without it.
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path

REQUIRED_COLUMNS = ("model", "input_per_1m", "output_per_1m")
OPTIONAL_COLUMNS = (
    "label",
    "cache_read_per_1m",
    "cache_write_per_1m",
    "currency",
    "discount_factor",
)


@dataclass(frozen=True, slots=True)
class ModelPrice:
    model: str
    label: str
    input_per_1m: float
    output_per_1m: float
    cache_read_per_1m: float | None
    cache_write_per_1m: float | None
    currency: str
    discount_factor: float | None = None


@dataclass(frozen=True, slots=True)
class PriceProblem:
    sheet_row: int
    reason: str


@dataclass(frozen=True, slots=True)
class PriceTable:
    prices: tuple[ModelPrice, ...]
    problems: tuple[PriceProblem, ...]

    def get(self, model: str) -> ModelPrice | None:
        return next((price for price in self.prices if price.model == model), None)


def load_price_table(path: Path, sheet_name: str | None = None) -> PriceTable:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name is not None else workbook.worksheets[0]
        rows = tuple(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    return parse_price_rows(rows)


def parse_price_rows(rows: Sequence[Sequence[object]]) -> PriceTable:
    if not rows:
        return PriceTable((), (PriceProblem(0, "price sheet is empty"),))

    header = tuple(_normalise_header(cell) for cell in rows[0])
    missing = tuple(column for column in REQUIRED_COLUMNS if column not in header)
    if missing:
        return PriceTable(
            (),
            (PriceProblem(1, f"missing required column(s): {', '.join(missing)}; found: {', '.join(header)}"),),
        )

    parsed = tuple(
        _parse_price_row(index, header, cells)
        for index, cells in enumerate(rows[1:], start=2)
        if any(cell is not None and str(cell).strip() != "" for cell in cells)
    )
    return PriceTable(
        prices=tuple(item for item in parsed if isinstance(item, ModelPrice)),
        problems=tuple(item for item in parsed if isinstance(item, PriceProblem)),
    )


def _parse_price_row(sheet_row: int, header: Sequence[str], cells: Sequence[object]) -> ModelPrice | PriceProblem:
    values = {name: cells[index] for index, name in enumerate(header) if index < len(cells) and name}

    model = str(values.get("model") or "").strip()
    if not model:
        return PriceProblem(sheet_row, "'model' is empty")

    input_per_1m = _parse_money(values.get("input_per_1m"))
    output_per_1m = _parse_money(values.get("output_per_1m"))
    if input_per_1m is None or output_per_1m is None:
        return PriceProblem(sheet_row, f"'{model}' has a non-numeric input_per_1m/output_per_1m")

    return ModelPrice(
        model=model,
        label=str(values.get("label") or "").strip() or model,
        input_per_1m=input_per_1m,
        output_per_1m=output_per_1m,
        cache_read_per_1m=_parse_money(values.get("cache_read_per_1m")),
        cache_write_per_1m=_parse_money(values.get("cache_write_per_1m")),
        currency=str(values.get("currency") or "").strip().upper() or "USD",
        discount_factor=_parse_discount_factor(values.get("discount_factor")),
    )


def _normalise_header(cell: object) -> str:
    return str(cell or "").strip().lower().replace(" ", "_")


def _parse_money(value: object) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().lstrip("$￥¥€£").replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _parse_discount_factor(value: object) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None
