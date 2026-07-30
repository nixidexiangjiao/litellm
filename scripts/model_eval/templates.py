"""Emit starter workbooks so the input format never has to be guessed.

They double as the demo inputs: the workload has two users out of chronological
order so the replay ordering is visible, and the price sheet leads with the
bundled mock provider so a fresh sandbox can produce a real report with no
vendor key at all.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import TYPE_CHECKING

from .pricing import OPTIONAL_COLUMNS, REQUIRED_COLUMNS

if TYPE_CHECKING:
    from openpyxl import Workbook

_PROMPTS = (
    ("u-1001", "2026-07-30 09:15:00", "Summarise the difference between TTFT and TPOT in two sentences"),
    ("u-1002", "2026-07-30 09:14:00", "Write a haiku about cache hit rates"),
    ("u-1001", "2026-07-30 09:16:30", "Now give me one concrete example of each"),
)

_PRICE_ROWS = (
    ("mock-fast", 0.27, 1.10, "bundled mock provider", 0.07, 0.0, "USD"),
    ("deepseek-official", 0.27, 1.10, "DeepSeek / deepseek-chat", 0.07, 0.0, "USD"),
    ("deepseek-siliconflow", 2.0, 8.0, "SiliconFlow / DeepSeek-V3", 0.4, 0.0, "CNY"),
)


def write_workload_template(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "requests"
    sheet.append(["user_id", "timestamp", "body"])
    for user_id, timestamp, prompt in _PROMPTS:
        sheet.append([user_id, timestamp, json.dumps(_body(prompt), ensure_ascii=False)])
    _finish(workbook, sheet_widths=(16, 22, 120), path=path)


def write_pricing_template(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "pricing"
    sheet.append([*REQUIRED_COLUMNS, *OPTIONAL_COLUMNS])
    for row in _PRICE_ROWS:
        sheet.append(list(row))
    _finish(workbook, sheet_widths=(24, 14, 14, 28, 20, 20, 10), path=path)


def _body(prompt: str) -> dict[str, object]:
    return {
        "model": "recorded-model-name",
        "stream": True,
        "messages": [{"role": "user", "content": prompt}],
    }


def _finish(workbook: "Workbook", sheet_widths: tuple[int, ...], path: Path) -> None:
    from openpyxl.styles import Font

    sheet = workbook.worksheets[0]
    for index, width in enumerate(sheet_widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    workbook.save(path)
