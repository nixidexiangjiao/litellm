"""Load the replay workload from an Excel sheet.

Column layout is positional, matching how the traffic export is produced:
column A is the request id, column B is the request StartTime, column C is the
UserId, and column D is the recorded OpenAI ``/v1/chat/completions`` request body as JSON text.

``openpyxl`` is imported inside the loader rather than at module scope: it is
the one dependency this tool adds on top of LiteLLM's own, and every parsing
and ordering rule below stays importable (and testable) without it.
"""

from __future__ import annotations

import json
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime, timezone
from numbers import Real
from pathlib import Path
from typing import Union

JSONValue = Union[None, bool, int, float, str, "list[JSONValue]", "dict[str, JSONValue]"]
JSONObject = dict[str, JSONValue]

_EPOCH_MILLIS_THRESHOLD = 10_000_000_000
_TIMESTAMP_FORMATS = (
    "%Y-%m-%d %H:%M:%S.%f",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d",
    "%Y/%m/%d %H:%M:%S",
    "%Y/%m/%d %H:%M",
    "%Y/%m/%d",
)


@dataclass(frozen=True, slots=True)
class WorkloadRow:
    sheet_row: int
    user_id: str
    timestamp: datetime
    body: JSONObject

    @property
    def recorded_model(self) -> str:
        recorded = self.body.get("model")
        return recorded if isinstance(recorded, str) else ""

    @property
    def recorded_stream(self) -> bool:
        return self.body.get("stream") is True


@dataclass(frozen=True, slots=True)
class RowProblem:
    sheet_row: int
    reason: str


@dataclass(frozen=True, slots=True)
class Workload:
    rows: tuple[WorkloadRow, ...]
    problems: tuple[RowProblem, ...]


def load_workload(path: Path, sheet_name: str | None = None) -> Workload:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name is not None else workbook.worksheets[0]
        raw_rows = tuple(
            (index, cells)
            for index, cells in enumerate(worksheet.iter_rows(values_only=True), start=1)
            if _has_content(cells)
        )
    finally:
        workbook.close()

    return _build_workload(raw_rows[1:] if _is_header(raw_rows) else raw_rows)


def order_rows(rows: Sequence[WorkloadRow]) -> tuple[WorkloadRow, ...]:
    """Order by user id, then by timestamp, keeping the sheet order as tie-break."""
    return tuple(sorted(rows, key=lambda row: (row.user_id, row.timestamp, row.sheet_row)))


def _build_workload(raw_rows: Sequence[tuple[int, Sequence[object]]]) -> Workload:
    parsed = tuple(_parse_row(index, cells) for index, cells in raw_rows)
    return Workload(
        rows=tuple(item for item in parsed if isinstance(item, WorkloadRow)),
        problems=tuple(item for item in parsed if isinstance(item, RowProblem)),
    )


def _parse_row(sheet_row: int, cells: Sequence[object]) -> WorkloadRow | RowProblem:
    if len(cells) < 4:
        return RowProblem(sheet_row, f"expected 4 columns, found {len(cells)}")

    user_id = _parse_user_id(cells[2])
    if user_id is None:
        return RowProblem(sheet_row, "column C (UserId) is empty")

    timestamp = _parse_timestamp(cells[1])
    if timestamp is None:
        return RowProblem(sheet_row, f"column B (StartTime) is not a recognisable timestamp: {cells[1]!r}")

    body = _parse_body(cells[3])
    if body is None:
        return RowProblem(sheet_row, "column D (Input) is not a JSON object")
    if not isinstance(body.get("messages"), list):
        return RowProblem(sheet_row, "column D (Input) has no 'messages' array")

    return WorkloadRow(sheet_row=sheet_row, user_id=user_id, timestamp=timestamp, body=body)


def _has_content(cells: Sequence[object]) -> bool:
    return any(cell is not None and str(cell).strip() != "" for cell in cells)


def _is_header(raw_rows: Sequence[tuple[int, Sequence[object]]]) -> bool:
    if not raw_rows:
        return False
    _, cells = raw_rows[0]
    return len(cells) < 4 or _parse_body(cells[3]) is None


def _parse_user_id(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def _parse_timestamp(value: object) -> datetime | None:
    if isinstance(value, datetime):
        return value if value.tzinfo is not None else value.replace(tzinfo=timezone.utc)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=timezone.utc)
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return _from_epoch(float(value))
    if isinstance(value, str):
        return _parse_timestamp_text(value.strip())
    return None


def _parse_timestamp_text(text: str) -> datetime | None:
    if not text:
        return None
    if text.replace(".", "", 1).isdigit():
        return _from_epoch(float(text))
    parsed = _try_fromisoformat(text)
    return parsed if parsed is not None else _try_strptime(text)


def _try_fromisoformat(text: str) -> datetime | None:
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return None
    return parsed if parsed.tzinfo is not None else parsed.replace(tzinfo=timezone.utc)


def _try_strptime(text: str) -> datetime | None:
    for fmt in _TIMESTAMP_FORMATS:
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    return None


def _from_epoch(value: float) -> datetime | None:
    seconds = value / 1000 if value >= _EPOCH_MILLIS_THRESHOLD else value
    try:
        return datetime.fromtimestamp(seconds, tz=timezone.utc)
    except (OverflowError, OSError, ValueError):
        return None


def _parse_body(value: object) -> JSONObject | None:
    if isinstance(value, Mapping):
        return {str(key): _as_json(item) for key, item in value.items()}
    if not isinstance(value, str):
        return None
    try:
        decoded = json.loads(value)
    except (TypeError, ValueError):
        return None
    return decoded if isinstance(decoded, dict) else None


def _as_json(value: object) -> JSONValue:
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, Mapping):
        return {str(key): _as_json(item) for key, item in value.items()}
    if isinstance(value, Sequence):
        return [_as_json(item) for item in value]
    return str(value)
