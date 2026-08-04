"""Present what the database already computed: a console table and an xlsx dump.

Nothing is aggregated here. The numbers come from the ``eval_summary`` and
``eval_request_costs`` views so the spreadsheet and the terminal can never
disagree with a query someone writes later.

``openpyxl`` is imported inside the exporter; the console table needs no
dependency at all.
"""

from __future__ import annotations

from collections.abc import Sequence
from datetime import datetime
from pathlib import Path

from .storage import EvalStore, SummaryRow

EXCEL_CELL_LIMIT = 32_767

_CONSOLE_HEADERS = (
    "model",
    "ok/total",
    "ttft p50",
    "tpot p50",
    "total p50",
    "out tok/s",
    "cache hit",
    "cost",
    "cost/req",
    "cost/1M out (all-in)",
)


def format_console_table(summaries: Sequence[SummaryRow]) -> str:
    rows = tuple(_console_row(summary) for summary in summaries)
    widths = tuple(max(len(cell) for cell in column) for column in zip(_CONSOLE_HEADERS, *rows, strict=True))
    return "\n".join(
        " | ".join(cell.ljust(width) for cell, width in zip(row, widths, strict=True))
        for row in (_CONSOLE_HEADERS, *rows)
    )


def export_workbook(store: EvalStore, path: Path, run_id: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    for title, query in (
        ("summary", "SELECT * FROM eval_summary WHERE run_id = ? ORDER BY model"),
        (
            "requests",
            "SELECT * FROM eval_request_costs WHERE run_id = ? ORDER BY model, repetition, sequence",
        ),
    ):
        cursor = store.connection.execute(query, [run_id])
        headers = tuple(str(column[0]) for column in cursor.description or ())
        sheet = workbook.create_sheet(title)
        sheet.append(list(headers))
        for row in cursor.fetchall():
            sheet.append([_to_cell(value) for value in row])
        for index, header in enumerate(headers, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 2, 10), 40)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
    workbook.save(path)


def _console_row(summary: SummaryRow) -> tuple[str, ...]:
    currency = summary.currency or "?"
    return (
        summary.label or summary.model,
        f"{summary.succeeded}/{summary.requests}",
        _fmt(summary.ttft_ms_p50),
        _fmt(summary.tpot_ms_p50),
        _fmt(summary.total_ms_p50),
        _fmt(summary.output_tokens_per_s_mean),
        _fmt_ratio(summary.cache_hit_rate),
        f"{summary.cost_total:.4f} {currency}",
        _fmt(summary.cost_per_request, digits=6),
        _fmt(summary.cost_per_1m_output_tokens, digits=4),
    )


def _fmt(value: float | None, digits: int = 2) -> str:
    return "-" if value is None else f"{value:.{digits}f}"


def _fmt_ratio(value: float | None) -> str:
    return "-" if value is None else f"{value * 100:.1f}%"


def _to_cell(value: object) -> str | float | int | bool | None:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, datetime):
        return value.isoformat()
    text = str(value)
    return text if len(text) <= EXCEL_CELL_LIMIT else f"{text[: EXCEL_CELL_LIMIT - 3]}..."
